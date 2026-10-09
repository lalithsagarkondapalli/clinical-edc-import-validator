"""Readers: turn each approved source format into a list of SourceRecord.

Every reader returns values as *strings* (or None) so that all type checking
happens in one place (validate.py) regardless of the source format.
"""
from __future__ import annotations

import csv
import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook


@dataclass
class SourceRecord:
    record_no: int                 # 1-based position in the source (row / element)
    values: dict[str, str | None]


@dataclass
class FileIssue:
    rule_id: str
    message: str
    severity: str = "error"


@dataclass
class ReadResult:
    records: list[SourceRecord]
    sha256: str
    bytes: int
    issues: list[FileIssue] = field(default_factory=list)


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def read_csv(path: Path, cfg: dict) -> tuple[list[SourceRecord], list[FileIssue]]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        recs = [SourceRecord(i, {k: _clean(v) for k, v in row.items()})
                for i, row in enumerate(reader, start=1)]
    return recs, []


def read_xlsx(path: Path, cfg: dict) -> tuple[list[SourceRecord], list[FileIssue]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[cfg.get("sheet") or wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    recs = []
    for i, row in enumerate(rows, start=1):
        if row is None or all(c in (None, "") for c in row):
            continue
        recs.append(SourceRecord(i, {h: _clean(v) for h, v in zip(header, row)}))
    wb.close()
    return recs, []


def read_json(path: Path, cfg: dict) -> tuple[list[SourceRecord], list[FileIssue]]:
    doc = json.loads(path.read_text(encoding="utf-8"))
    rows = doc[cfg.get("records_path", "records")]
    issues: list[FileIssue] = []
    header = doc.get("header", {})
    declared = header.get("record_count")
    if declared is not None and int(declared) != len(rows):
        issues.append(FileIssue("F-COUNT",
                                f"header record_count={declared} but file contains {len(rows)} records"))
    recs = [SourceRecord(i, {k: _clean(v) for k, v in r.items()}) for i, r in enumerate(rows, start=1)]
    return recs, issues


def read_xml(path: Path, cfg: dict) -> tuple[list[SourceRecord], list[FileIssue]]:
    parser = etree.XMLParser(resolve_entities=False, no_network=True)  # XXE-safe
    tree = etree.parse(str(path), parser)
    issues: list[FileIssue] = []
    if cfg.get("schema"):
        schema = etree.XMLSchema(etree.parse(cfg["_schema_path"]))
        if not schema.validate(tree):
            for err in schema.error_log:
                issues.append(FileIssue("F-XSD", f"line {err.line}: {err.message}"))
    ns = cfg.get("namespaces", {})
    recs = []
    for i, node in enumerate(tree.xpath(cfg["record_xpath"], namespaces=ns), start=1):
        values = {}
        for name, xp in cfg["fields"].items():
            hit = node.xpath(xp, namespaces=ns)
            if not hit:
                values[name] = None
            else:
                h = hit[0]
                values[name] = _clean(h if isinstance(h, str) else h.text)
        recs.append(SourceRecord(i, values))
    return recs, issues


READERS = {"csv": read_csv, "xlsx": read_xlsx, "json": read_json, "xml": read_xml}


def read_source(source_cfg: dict, base_dir: Path) -> ReadResult:
    path = (base_dir / source_cfg["path"]).resolve()
    cfg = dict(source_cfg)
    if cfg.get("schema"):
        cfg["_schema_path"] = str(base_dir / cfg["schema"])
    fmt = cfg["format"].lower()
    if fmt not in READERS:
        raise ValueError(f"Unsupported source format: {fmt}")
    records, issues = READERS[fmt](path, cfg)
    return ReadResult(records, sha256_of(path), path.stat().st_size, issues)
